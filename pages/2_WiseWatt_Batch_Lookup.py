from __future__ import annotations

import io
import re
import time
from pathlib import Path
from typing import Any

import requests
import streamlit as st
from openpyxl import load_workbook

try:
    import truststore

    truststore.inject_into_ssl()
except ImportError:
    pass


# ============================================================
# STREAMLIT PAGE
# ============================================================

st.set_page_config(
    page_title="WiseWatt Batch Lookup",
    page_icon="⚡",
    layout="wide",
)

st.title("⚡ WiseWatt MPAN / MPRN Batch Lookup")
st.caption(
    "Upload an Excel workbook, retrieve electricity MPANs and gas MPRNs, "
    "and download the updated workbook."
)


# ============================================================
# CONFIGURATION
# ============================================================

DEFAULT_API_URL = "https://api.wisewatt.uk.cgi.com/api/v1"
DEFAULT_REQUEST_TIMEOUT = 45
DEFAULT_MAX_RETRIES = 3
DEFAULT_REQUEST_DELAY = 0.25

ADDRESS_COLUMN = "ADDRESS"
ELECTRIC_METER_COLUMN = "ELECTRIC METER"
GAS_METER_COLUMN = "GAS METER"

OUTPUT_COLUMNS = [
    "EXTRACTED ADDRESS IDENTIFIER",
    "EXTRACTED POSTCODE",
    "ELECTRICITY MPAN",
    "ELECTRICITY STATUS",
    "ELECTRICITY ERROR CODE",
    "ELECTRICITY TRANSACTION ID",
    "GAS MPRN",
    "GAS STATUS",
    "GAS ERROR CODE",
    "GAS TRANSACTION ID",
]

ERROR_DESCRIPTIONS = {
    "E080201": "The request does not uniquely identify a premises",
    "E080202": "The premises do not contain any devices",
}


def read_secret(name: str, default: str = "") -> str:
    try:
        return str(st.secrets.get(name, default)).strip()
    except Exception:
        return default


API_URL = read_secret("WISEWATT_API_URL", DEFAULT_API_URL).rstrip("/")
SECRET_API_KEY = read_secret("WISEWATT_API_KEY")
INVENTORY_ENDPOINT = f"{API_URL}/inventory/manual"


# ============================================================
# GENERAL HELPERS
# ============================================================

def is_blank(value: Any) -> bool:
    """Return True for None, empty strings and whitespace-only strings."""
    return value is None or str(value).strip() == ""


def meter_required(value: Any) -> bool:
    """
    Interpret the ELECTRIC METER / GAS METER flag.

    Values treated as false:
      blank, No, N, False, 0, None, N/A, NA

    Any other non-empty value is treated as requesting a lookup.
    """
    if is_blank(value):
        return False

    return str(value).strip().lower() not in {
        "no",
        "n",
        "false",
        "0",
        "none",
        "n/a",
        "na",
    }


# ============================================================
# ADDRESS PARSING
# ============================================================

UK_POSTCODE_PATTERN = re.compile(
    r"\b("
    r"GIR\s?0AA|"
    r"(?:"
    r"[A-PR-UWYZ][0-9][0-9A-HJKSTUW]?|"
    r"[A-PR-UWYZ][A-HK-Y][0-9][0-9ABEHMNPRV-Y]?"
    r")"
    r"\s?[0-9][ABD-HJLNP-UW-Z]{2}"
    r")\b",
    flags=re.IGNORECASE,
)


def normalise_postcode(postcode: str) -> str:
    compact = re.sub(r"\s+", "", postcode.strip().upper())
    if len(compact) < 5:
        return compact
    return f"{compact[:-3]} {compact[-3:]}"


def parse_address(
    full_address: Any,
) -> tuple[str | None, str | None, str | None]:
    """
    Extract the DCC AddressIdentifier and UK postcode.

    Expected example:
        29, Calluna Close, Dowlais, Merthyr Tydfil, CF48 3HF

    The first comma-separated item becomes the DCC AddressIdentifier.
    This may be a house number, suffix, flat identifier or house name.
    """
    if is_blank(full_address):
        return None, None, "Address is empty"

    address = str(full_address).strip()
    parts = [part.strip() for part in address.split(",") if part.strip()]

    if len(parts) < 2:
        return None, None, "Address is not comma-separated"

    address_identifier = parts[0]

    postcode_matches = list(UK_POSTCODE_PATTERN.finditer(address))
    if not postcode_matches:
        return address_identifier, None, "No valid UK postcode was found"

    postcode = normalise_postcode(postcode_matches[-1].group(1))

    if len(address_identifier) > 30:
        return (
            None,
            postcode,
            "Address identifier exceeds the 30-character DCC limit",
        )

    return address_identifier, postcode, None


# ============================================================
# EXCEL HELPERS
# ============================================================

def get_header_map(worksheet) -> dict[str, int]:
    """Build a case-insensitive map of header text to column number."""
    header_map: dict[str, int] = {}

    for cell in worksheet[1]:
        if not is_blank(cell.value):
            header_map[str(cell.value).strip().upper()] = cell.column

    return header_map


def ensure_output_columns(worksheet) -> dict[str, int]:
    """Add output headings if they are not already present."""
    header_map = get_header_map(worksheet)

    for heading in OUTPUT_COLUMNS:
        key = heading.upper()

        if key not in header_map:
            column_number = worksheet.max_column + 1
            worksheet.cell(
                row=1,
                column=column_number,
                value=heading,
            )
            header_map[key] = column_number

    return header_map


def find_rows_to_process(
    worksheet,
    address_column: int,
    electricity_flag_column: int,
    gas_flag_column: int,
) -> list[int]:
    """
    Return only meaningful spreadsheet rows.

    Completely empty rows are ignored even when worksheet.max_row is inflated
    by old formatting, borders, deleted values or previous Excel edits.

    A row is meaningful when at least one of these source fields contains data:
      ADDRESS, ELECTRIC METER, GAS METER
    """
    meaningful_rows: list[int] = []

    for row_number in range(2, worksheet.max_row + 1):
        values = (
            worksheet.cell(row=row_number, column=address_column).value,
            worksheet.cell(
                row=row_number,
                column=electricity_flag_column,
            ).value,
            worksheet.cell(row=row_number, column=gas_flag_column).value,
        )

        if any(not is_blank(value) for value in values):
            meaningful_rows.append(row_number)

    return meaningful_rows


def write_result(
    worksheet,
    row_number: int,
    meter_type: str,
    result: dict[str, Any],
    header_map: dict[str, int],
) -> None:
    if meter_type == "electricity":
        number_heading = "ELECTRICITY MPAN"
        status_heading = "ELECTRICITY STATUS"
        error_heading = "ELECTRICITY ERROR CODE"
        transaction_heading = "ELECTRICITY TRANSACTION ID"
    else:
        number_heading = "GAS MPRN"
        status_heading = "GAS STATUS"
        error_heading = "GAS ERROR CODE"
        transaction_heading = "GAS TRANSACTION ID"

    worksheet.cell(
        row=row_number,
        column=header_map[number_heading],
        value=result.get("meter_number", ""),
    )
    worksheet.cell(
        row=row_number,
        column=header_map[status_heading],
        value=result.get("status", ""),
    )
    worksheet.cell(
        row=row_number,
        column=header_map[error_heading],
        value=result.get("error_code", ""),
    )
    worksheet.cell(
        row=row_number,
        column=header_map[transaction_heading],
        value=result.get("transaction_id", ""),
    )


def existing_meter_number(
    worksheet,
    row_number: int,
    heading: str,
    header_map: dict[str, int],
) -> str:
    column = header_map.get(heading.upper())
    if not column:
        return ""

    value = worksheet.cell(row=row_number, column=column).value
    return "" if is_blank(value) else str(value).strip()


# ============================================================
# WISEWATT API
# ============================================================

def create_session(api_key: str) -> requests.Session:
    session = requests.Session()
    session.headers.update(
        {
            "wisewatt-api-key-token": api_key,
            "Accept": "application/json",
            "Content-Type": "application/json",
            "User-Agent": "Safehouse-WiseWatt-Batch-Lookup/2.0",
        }
    )
    return session


def decode_response(response: requests.Response) -> dict[str, Any]:
    try:
        payload = response.json()
    except ValueError:
        return {
            "_invalid_json": True,
            "_raw_text": response.text[:500],
        }

    if isinstance(payload, dict):
        return payload

    return {
        "_unexpected_json": True,
        "_json_value": payload,
    }


def lookup_meter(
    session: requests.Session,
    postcode: str,
    address_identifier: str,
    meter_type: str,
    request_timeout: int,
    max_retries: int,
) -> dict[str, Any]:
    payload = {
        "postCode": postcode,
        "addressIdentifier": address_identifier,
        "meterType": meter_type,
    }

    last_error = "Unknown request error"

    for attempt in range(1, max_retries + 1):
        try:
            response = session.post(
                INVENTORY_ENDPOINT,
                json=payload,
                timeout=request_timeout,
                verify=True,
            )

            data = decode_response(response)
            transaction_id = data.get("transactionId", "")
            error_code = data.get("errorCode")
            device = data.get("device")

            if response.ok and isinstance(device, dict):
                meter_number = device.get("importMpxn")

                if not is_blank(meter_number):
                    meter_label = (
                        "MPAN"
                        if meter_type == "electricity"
                        else "MPRN"
                    )

                    return {
                        "success": True,
                        "meter_number": str(meter_number),
                        "transaction_id": transaction_id,
                        "error_code": "",
                        "status": (
                            f"Success – {meter_label} found; "
                            f"device status: "
                            f"{device.get('deviceStatus', 'unknown')}"
                        ),
                    }

                return {
                    "success": False,
                    "meter_number": "",
                    "transaction_id": transaction_id,
                    "error_code": "MISSING_MPXN",
                    "status": (
                        "A device was returned, but importMpxn was missing"
                    ),
                }

            if error_code:
                error_code = str(error_code)
                description = ERROR_DESCRIPTIONS.get(
                    error_code,
                    "DCC/WiseWatt returned an application error",
                )

                return {
                    "success": False,
                    "meter_number": "",
                    "transaction_id": transaction_id,
                    "error_code": error_code,
                    "status": f"{error_code} – {description}",
                }

            detail = (
                data.get("message")
                or data.get("_raw_text")
                or str(data)
            )

            if 400 <= response.status_code < 500:
                return {
                    "success": False,
                    "meter_number": "",
                    "transaction_id": transaction_id,
                    "error_code": f"HTTP_{response.status_code}",
                    "status": (
                        f"HTTP {response.status_code} – {detail}"
                    ),
                }

            last_error = f"HTTP {response.status_code} – {detail}"

        except requests.exceptions.Timeout:
            last_error = (
                f"Request timed out after {request_timeout} seconds"
            )

        except requests.exceptions.SSLError as exc:
            return {
                "success": False,
                "meter_number": "",
                "transaction_id": "",
                "error_code": "SSL_ERROR",
                "status": f"SSL certificate error – {exc}",
            }

        except requests.exceptions.ConnectionError as exc:
            last_error = f"Connection error – {exc}"

        except requests.exceptions.RequestException as exc:
            last_error = f"Request error – {exc}"

        if attempt < max_retries:
            time.sleep(2 ** (attempt - 1))

    return {
        "success": False,
        "meter_number": "",
        "transaction_id": "",
        "error_code": "REQUEST_FAILED",
        "status": (
            f"Failed after {max_retries} attempts – {last_error}"
        ),
    }


# ============================================================
# SIDEBAR
# ============================================================

with st.sidebar:
    st.header("Configuration")

    st.text_input(
        "WiseWatt API URL",
        value=API_URL,
        disabled=True,
    )

    if SECRET_API_KEY:
        st.success("API key loaded from Streamlit secrets.")
        api_key = SECRET_API_KEY
    else:
        st.warning(
            "WISEWATT_API_KEY was not found in Streamlit secrets."
        )
        api_key = st.text_input(
            "Temporary API key",
            type="password",
            help="Used only for the current Streamlit session.",
        ).strip()

    request_delay = st.number_input(
        "Delay between API requests (seconds)",
        min_value=0.0,
        max_value=5.0,
        value=DEFAULT_REQUEST_DELAY,
        step=0.05,
    )

    request_timeout = st.number_input(
        "Request timeout (seconds)",
        min_value=5,
        max_value=180,
        value=DEFAULT_REQUEST_TIMEOUT,
        step=5,
    )

    max_retries = st.number_input(
        "Maximum attempts per request",
        min_value=1,
        max_value=5,
        value=DEFAULT_MAX_RETRIES,
        step=1,
    )

    skip_existing = st.checkbox(
        "Skip rows with an existing MPAN/MPRN",
        value=True,
    )

    st.markdown("**Required source columns**")
    st.code(
        "ADDRESS\nELECTRIC METER\nGAS METER",
        language=None,
    )


# ============================================================
# FILE UPLOAD
# ============================================================

uploaded_file = st.file_uploader(
    "Upload Excel workbook",
    type=["xlsx"],
    help=(
        "Example ADDRESS: "
        "29, Calluna Close, Dowlais, Merthyr Tydfil, CF48 3HF"
    ),
)

if uploaded_file is None:
    st.info("Upload an .xlsx workbook to begin.")
    st.stop()

try:
    original_workbook_bytes = uploaded_file.getvalue()
    preview_workbook = load_workbook(
        io.BytesIO(original_workbook_bytes)
    )
except Exception as exc:
    st.error(f"Could not open the workbook: {exc}")
    st.stop()

sheet_name = st.selectbox(
    "Worksheet to process",
    options=preview_workbook.sheetnames,
)

preview_worksheet = preview_workbook[sheet_name]
preview_header_map = get_header_map(preview_worksheet)

required_columns = {
    ADDRESS_COLUMN,
    ELECTRIC_METER_COLUMN,
    GAS_METER_COLUMN,
}

missing_columns = sorted(
    column
    for column in required_columns
    if column.upper() not in preview_header_map
)

if missing_columns:
    st.error(
        "Missing required columns: " + ", ".join(missing_columns)
    )
    st.stop()

preview_rows = find_rows_to_process(
    preview_worksheet,
    preview_header_map[ADDRESS_COLUMN],
    preview_header_map[ELECTRIC_METER_COLUMN],
    preview_header_map[GAS_METER_COLUMN],
)

ignored_empty_rows = (
    max(preview_worksheet.max_row - 1, 0) - len(preview_rows)
)

metric1, metric2, metric3, metric4 = st.columns(4)
metric1.metric(
    "Meaningful rows",
    len(preview_rows),
)
metric2.metric(
    "Empty/formatted rows ignored",
    max(ignored_empty_rows, 0),
)
metric3.metric(
    "Worksheet",
    sheet_name,
)
metric4.metric(
    "Detected final Excel row",
    preview_worksheet.max_row,
)

st.info(
    "Rows are processed only when ADDRESS, ELECTRIC METER or GAS METER "
    "contains a value. Completely empty rows are silently ignored, even "
    "when Excel reports them as part of the used range because of formatting."
)

run_lookup = st.button(
    "Run meter lookup",
    type="primary",
    disabled=not bool(api_key) or not preview_rows,
)

if not run_lookup:
    st.stop()


# ============================================================
# PROCESS WORKBOOK
# ============================================================

workbook = load_workbook(io.BytesIO(original_workbook_bytes))
worksheet = workbook[sheet_name]
header_map = ensure_output_columns(worksheet)

address_column = header_map[ADDRESS_COLUMN]
electricity_flag_column = header_map[ELECTRIC_METER_COLUMN]
gas_flag_column = header_map[GAS_METER_COLUMN]

rows_to_process = find_rows_to_process(
    worksheet,
    address_column,
    electricity_flag_column,
    gas_flag_column,
)

session = create_session(api_key)

progress_bar = st.progress(0.0)
current_status = st.empty()
live_log = st.empty()

logs: list[str] = []

electricity_successes = 0
gas_successes = 0
failed_lookups = 0
address_errors = 0
skipped_flags = 0
skipped_existing = 0
api_requests = 0

total_rows = len(rows_to_process)

for position, row_number in enumerate(rows_to_process, start=1):
    full_address = worksheet.cell(
        row=row_number,
        column=address_column,
    ).value

    electricity_flag = worksheet.cell(
        row=row_number,
        column=electricity_flag_column,
    ).value
    gas_flag = worksheet.cell(
        row=row_number,
        column=gas_flag_column,
    ).value

    electricity_required = meter_required(electricity_flag)
    gas_required = meter_required(gas_flag)

    existing_mpan = existing_meter_number(
        worksheet,
        row_number,
        "ELECTRICITY MPAN",
        header_map,
    )
    existing_mprn = existing_meter_number(
        worksheet,
        row_number,
        "GAS MPRN",
        header_map,
    )

    current_status.write(
        f"Processing meaningful row {position} of {total_rows} "
        f"(Excel row {row_number})"
    )

    address_identifier, postcode, parsing_error = parse_address(
        full_address
    )

    worksheet.cell(
        row=row_number,
        column=header_map["EXTRACTED ADDRESS IDENTIFIER"],
        value=address_identifier or "",
    )
    worksheet.cell(
        row=row_number,
        column=header_map["EXTRACTED POSTCODE"],
        value=postcode or "",
    )

    # A row with only flags but no address is a genuine data error.
    # Fully blank rows never reach this loop.
    if parsing_error:
        if electricity_required:
            write_result(
                worksheet,
                row_number,
                "electricity",
                {
                    "meter_number": existing_mpan,
                    "transaction_id": "",
                    "error_code": "ADDRESS_ERROR",
                    "status": f"ADDRESS_ERROR – {parsing_error}",
                },
                header_map,
            )
            failed_lookups += 1

        if gas_required:
            write_result(
                worksheet,
                row_number,
                "gas",
                {
                    "meter_number": existing_mprn,
                    "transaction_id": "",
                    "error_code": "ADDRESS_ERROR",
                    "status": f"ADDRESS_ERROR – {parsing_error}",
                },
                header_map,
            )
            failed_lookups += 1

        if electricity_required or gas_required:
            address_errors += 1
            logs.append(
                f"Excel row {row_number}: ADDRESS_ERROR – "
                f"{parsing_error}"
            )
        else:
            logs.append(
                f"Excel row {row_number}: skipped – no meter lookup flags"
            )
            skipped_flags += 2

        progress_bar.progress(position / total_rows)
        live_log.code("\n".join(logs[-15:]), language=None)
        continue

    if electricity_required:
        if skip_existing and existing_mpan:
            skipped_existing += 1
            write_result(
                worksheet,
                row_number,
                "electricity",
                {
                    "meter_number": existing_mpan,
                    "transaction_id": "",
                    "error_code": "",
                    "status": "Skipped – existing MPAN retained",
                },
                header_map,
            )
            logs.append(
                f"Excel row {row_number}: existing MPAN retained"
            )
        else:
            result = lookup_meter(
                session=session,
                postcode=postcode,
                address_identifier=address_identifier,
                meter_type="electricity",
                request_timeout=int(request_timeout),
                max_retries=int(max_retries),
            )
            api_requests += 1
            write_result(
                worksheet,
                row_number,
                "electricity",
                result,
                header_map,
            )

            if result["success"]:
                electricity_successes += 1
                logs.append(
                    f"Excel row {row_number}: MPAN "
                    f"{result['meter_number']}"
                )
            else:
                failed_lookups += 1
                logs.append(
                    f"Excel row {row_number}: electricity – "
                    f"{result['status']}"
                )

            time.sleep(float(request_delay))
    else:
        skipped_flags += 1
        write_result(
            worksheet,
            row_number,
            "electricity",
            {
                "meter_number": existing_mpan,
                "transaction_id": "",
                "error_code": "",
                "status": "Skipped – ELECTRIC METER is No/blank",
            },
            header_map,
        )

    if gas_required:
        if skip_existing and existing_mprn:
            skipped_existing += 1
            write_result(
                worksheet,
                row_number,
                "gas",
                {
                    "meter_number": existing_mprn,
                    "transaction_id": "",
                    "error_code": "",
                    "status": "Skipped – existing MPRN retained",
                },
                header_map,
            )
            logs.append(
                f"Excel row {row_number}: existing MPRN retained"
            )
        else:
            result = lookup_meter(
                session=session,
                postcode=postcode,
                address_identifier=address_identifier,
                meter_type="gas",
                request_timeout=int(request_timeout),
                max_retries=int(max_retries),
            )
            api_requests += 1
            write_result(
                worksheet,
                row_number,
                "gas",
                result,
                header_map,
            )

            if result["success"]:
                gas_successes += 1
                logs.append(
                    f"Excel row {row_number}: MPRN "
                    f"{result['meter_number']}"
                )
            else:
                failed_lookups += 1
                logs.append(
                    f"Excel row {row_number}: gas – "
                    f"{result['status']}"
                )

            time.sleep(float(request_delay))
    else:
        skipped_flags += 1
        write_result(
            worksheet,
            row_number,
            "gas",
            {
                "meter_number": existing_mprn,
                "transaction_id": "",
                "error_code": "",
                "status": "Skipped – GAS METER is No/blank",
            },
            header_map,
        )

    progress_bar.progress(position / total_rows)
    live_log.code("\n".join(logs[-15:]), language=None)


# ============================================================
# SAVE AND DOWNLOAD
# ============================================================

output_buffer = io.BytesIO()
workbook.save(output_buffer)
output_bytes = output_buffer.getvalue()

output_name = (
    f"{Path(uploaded_file.name).stem}_with_meter_numbers.xlsx"
)

progress_bar.progress(1.0)
current_status.success("Processing complete.")

st.subheader("Results")

result1, result2, result3, result4, result5 = st.columns(5)
result1.metric("Electricity MPANs found", electricity_successes)
result2.metric("Gas MPRNs found", gas_successes)
result3.metric("Failed lookups", failed_lookups)
result4.metric("Address-error rows", address_errors)
result5.metric("API requests made", api_requests)

detail1, detail2, detail3 = st.columns(3)
detail1.metric("Empty rows ignored", max(ignored_empty_rows, 0))
detail2.metric("Meter flags skipped", skipped_flags)
detail3.metric("Existing numbers retained", skipped_existing)

st.download_button(
    "Download updated workbook",
    data=output_bytes,
    file_name=output_name,
    mime=(
        "application/vnd.openxmlformats-officedocument."
        "spreadsheetml.sheet"
    ),
    type="primary",
)

if failed_lookups:
    st.warning(
        "Some requested lookups failed. Review the status, error-code "
        "and transaction-ID columns in the downloaded workbook."
    )
else:
    st.success("All requested lookups completed without recorded failures.")
